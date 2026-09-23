from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

from docx import Document


def load_payload(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def format_deadline(task: Dict[str, Any]) -> str:
    if task.get("deadline"):
        return task["deadline"]
    if task.get("deadline_source"):
        return task["deadline_source"]
    return "— не определён —"


def format_review(task: Dict[str, Any]) -> str:
    return "Да" if bool(task.get("needs_review")) else ""


def source_label(task: Dict[str, Any]) -> str:
    speakers = task.get("source_speakers") or []
    if speakers:
        return ", ".join(speakers)
    return "UNKNOWN"


def add_tasks_table(document: Document, tasks: List[Dict[str, Any]]) -> None:
    table = document.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    header[0].text = "№"
    header[1].text = "Поручение"
    header[2].text = "Ответственный"
    header[3].text = "Срок"
    header[4].text = "Проверить"

    for index, task in enumerate(tasks, start=1):
        row = table.add_row().cells
        row[0].text = str(index)
        row[1].text = task.get("description") or ""
        row[2].text = task.get("assignee") or "— не определён —"
        row[3].text = format_deadline(task)
        row[4].text = format_review(task)


def add_sources(document: Document, tasks: List[Dict[str, Any]]) -> None:
    document.add_heading("Источники (для проверки)", level=2)
    for index, task in enumerate(tasks, start=1):
        paragraph = document.add_paragraph()
        label = f"[{index}] [{task.get('source_start', 0):.1f}s] [{source_label(task)}] "
        paragraph.add_run(label).bold = True
        paragraph.add_run(task.get("source_text") or "")


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python export_docx.py <result.json> [meeting_title]")

    input_path = Path(sys.argv[1])
    meeting_title = sys.argv[2] if len(sys.argv) > 2 else "Совещание"
    payload = load_payload(input_path)
    summary = payload.get("summary", "")
    tasks = payload.get("tasks", [])

    document = Document()
    document.add_heading(f"Протокол: {meeting_title}", level=1)
    document.add_paragraph(f"Дата формирования: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    document.add_paragraph("Автоматически сформировано системой DARAI. Требуется проверка секретарём.")

    document.add_heading("Саммари", level=2)
    document.add_paragraph(summary)

    document.add_heading("Поручения", level=2)
    add_tasks_table(document, tasks)
    add_sources(document, tasks)

    output_path = input_path.with_name(f"{input_path.stem}_protocol.docx")
    document.save(output_path)
    print(f"Saved to {output_path}")


if __name__ == "__main__":
    main()