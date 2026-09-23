"""DOCX/PDF protocol export from saved (secretary-edited) data. No LLM calls."""

from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .config import get_settings
from .errors import ApiError


@dataclass
class ExportTask:
    number: int
    task: str
    author: str
    assignee: str
    deadline: date | None
    deadline_source: str | None
    evidence: str | None
    execution_status: str


@dataclass
class ExportLine:
    start: float
    speaker: str
    text: str


@dataclass
class ExportData:
    title: str
    starts_at: datetime
    timezone: str
    agenda: str
    participants: list[tuple[str, str]]  # (fio, position)
    summary: str | None
    tasks: list[ExportTask]
    transcript: list[ExportLine] | None
    draft: bool
    protocol_version: int
    confirmed_at: datetime | None


def _set_font(document: Document, font: str) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Table Grid"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.name = font
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), font)


_TASK_COLUMNS_CM = (0.8, 4.2, 2.9, 2.7, 2.6, 3.3)  # 16.5 cm fits A4 with 2 cm margins


def _ts(seconds: float) -> str:
    s = int(seconds)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def build_docx(data: ExportData) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    _set_font(doc, get_settings().export_font)
    doc.styles["Normal"].font.size = Pt(11)

    if data.draft:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("ЧЕРНОВИК — протокол не подтверждён")
        run.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    doc.add_heading(f"Протокол совещания: {data.title}", level=1)
    local = data.starts_at.astimezone(ZoneInfo(data.timezone))
    doc.add_paragraph(f"Дата и время: {local.strftime('%d.%m.%Y %H:%M')} ({data.timezone})")
    if not data.draft and data.confirmed_at is not None:
        confirmed = data.confirmed_at.astimezone(ZoneInfo(data.timezone))
        doc.add_paragraph(f"Протокол подтверждён: {confirmed.strftime('%d.%m.%Y %H:%M')}, версия {data.protocol_version}")

    if data.agenda.strip():
        doc.add_heading("Повестка", level=2)
        for line in data.agenda.strip().splitlines():
            doc.add_paragraph(line)

    doc.add_heading("Участники", level=2)
    if data.participants:
        for fio, position in data.participants:
            doc.add_paragraph(f"{fio} — {position}", style="List Bullet")
    else:
        doc.add_paragraph("—")

    doc.add_heading("Краткое содержание", level=2)
    for block in (data.summary or "Саммари отсутствует.").split("\n"):
        doc.add_paragraph(block)

    doc.add_heading("Поручения", level=2)
    if data.tasks:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["№", "Поручение", "Исполнитель", "Срок", "Автор", "Основание"]):
            cell.text = ""
            cell.paragraphs[0].add_run(text).bold = True
        for t in data.tasks:
            deadline = t.deadline.strftime("%d.%m.%Y") if t.deadline else "не указан"
            if t.deadline_source:
                deadline += f" («{t.deadline_source}»)"
            row = table.add_row().cells
            for cell, text in zip(row, [str(t.number), t.task, t.assignee, deadline, t.author, t.evidence or "—"]):
                cell.text = text
        table.autofit = False
        for column, width in zip(table.columns, _TASK_COLUMNS_CM):
            column.width = Cm(width)  # tblGrid: honoured by LibreOffice
        for row in table.rows:
            for cell, width in zip(row.cells, _TASK_COLUMNS_CM):
                cell.width = Cm(width)
    else:
        doc.add_paragraph("Поручений нет.")

    if data.transcript is not None:
        doc.add_heading("Приложение: транскрипт", level=2)
        for line in data.transcript:
            p = doc.add_paragraph()
            p.add_run(f"[{_ts(line.start)}] {line.speaker}: ").bold = True
            p.add_run(line.text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def soffice_available() -> bool:
    return shutil.which(get_settings().soffice_path) is not None


def docx_to_pdf(docx_bytes: bytes) -> bytes:
    settings = get_settings()
    binary = shutil.which(settings.soffice_path)
    if binary is None:
        raise ApiError(503, "EXPORT_UNAVAILABLE", "LibreOffice не установлен: экспорт в PDF недоступен")
    settings.tmp_dir.mkdir(parents=True, exist_ok=True)
    # Isolated working dir and LibreOffice profile per conversion.
    with tempfile.TemporaryDirectory(prefix="export-", dir=settings.tmp_dir) as tmp:
        tmp_path = Path(tmp)
        src = tmp_path / "protocol.docx"
        src.write_bytes(docx_bytes)
        profile = (tmp_path / "profile").resolve().as_uri()
        try:
            proc = subprocess.run(
                [binary, f"-env:UserInstallation={profile}", "--headless", "--norestore",
                 "--convert-to", "pdf", "--outdir", str(tmp_path), str(src)],
                capture_output=True, timeout=settings.soffice_timeout_seconds, check=False,
                env={"HOME": str(tmp_path), "PATH": "/usr/bin:/bin:/usr/local/bin:/opt/homebrew/bin"},
            )
        except subprocess.TimeoutExpired:
            raise ApiError(503, "EXPORT_UNAVAILABLE", "LibreOffice превысил таймаут конвертации")
        out = tmp_path / "protocol.pdf"
        if proc.returncode != 0 or not out.is_file():
            raise ApiError(503, "EXPORT_UNAVAILABLE", "LibreOffice не смог сконвертировать документ")
        return out.read_bytes()
