import io
import shutil
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document
from sqlalchemy import func, select

from app.db import db_session
from app.models import Notification
from app.notifications import run_reminders

from .conftest import create_meeting, upload


def _processed_meeting(org, ai):
    sec = org["api"]["sec"]
    m = create_meeting(sec, [org["dana"]["id"], org["bolat"]["id"], org["sec"]["id"]])
    ai.speaker_map = {"SPEAKER_00": uuid.UUID(org["sec"]["id"]), "SPEAKER_01": uuid.UUID(org["dana"]["id"]),
                      "SPEAKER_02": None}
    upload(sec, m["id"])
    return sec, sec.get(f"/meetings/{m['id']}").json()


def _count(event: str) -> int:
    with db_session() as db:
        return db.scalar(select(func.count()).select_from(Notification).where(Notification.event_type == event))


def test_manual_edits_persist_and_speaker_rebinding_updates_author(org, ai):
    sec, d = _processed_meeting(org, ai)
    mid = d["id"]
    smeta = next(t for t in d["tasks"] if t["task"] == "Проверить смету")
    assert smeta["from"] == "SPEAKER_02"
    # unknown speaker identified manually -> author follows, assignee does not change
    r = sec.patch(f"/meetings/{mid}/speakers", json={"SPEAKER_02": org["bolat"]["id"]})
    assert r.status_code == 200 and r.json()["draft_revision"] == d["draft_revision"] + 1
    sp = {s["label"]: s for s in r.json()["speakers"]}["SPEAKER_02"]
    assert sp["confirmed_employee_id"] == org["bolat"]["id"] and sp["manually_set"] and sp["proposed_employee_id"] is None
    # non-participant and unknown label are rejected
    assert sec.patch(f"/meetings/{mid}/speakers", json={"SPEAKER_02": org["guest"]["id"]}).status_code == 422
    assert sec.patch(f"/meetings/{mid}/speakers", json={"SPEAKER_99": None}).status_code == 422
    # task edit
    r = sec.patch(f"/tasks/{smeta['id']}", json={"to": org["bolat"]["id"], "deadline": "2026-09-30",
                                                 "task": "Проверить смету проекта"})
    assert r.status_code == 200, r.text
    t = r.json()
    assert t["from"] == org["bolat"]["id"] and t["to"] == org["bolat"]["id"]
    assert "missing_assignee" not in t["review_reasons"] and "missing_deadline" not in t["review_reasons"]
    assert sec.patch(f"/tasks/{smeta['id']}", json={"to": org["guest"]["id"]}).status_code == 422
    # summary + manual task
    assert sec.patch(f"/meetings/{mid}/summary", json={"summary": "Итог: смета и отчёт (қазақша: есеп)."}).status_code == 200
    r = sec.post(f"/meetings/{mid}/tasks", json={"task": "Разослать протокол", "to": org["dana"]["id"],
                                                 "from": org["sec"]["id"], "deadline": "2026-09-24"})
    assert r.status_code == 201 and r.json()["origin"] == "manual" and r.json()["confidence"] is None
    # reload: all edits persisted
    d2 = sec.get(f"/meetings/{mid}").json()
    assert d2["summary"].startswith("Итог") and d2["summary_edited"] is True
    assert {t["task"] for t in d2["tasks"]} == {"Подготовить отчёт", "Проверить смету проекта", "Разослать протокол"}
    # participant cannot edit tasks
    assert org["api"]["dana"].patch(f"/tasks/{smeta['id']}", json={"task": "x"}).status_code == 404
    # draft task: assignee cannot mark execution
    report = next(t for t in d2["tasks"] if t["task"] == "Подготовить отчёт")
    assert sec.patch(f"/tasks/{report['id']}", json={"execution_status": "completed"}).json()["error"]["code"] == \
        "TASK_NOT_CONFIRMED"
    assert sec.patch(f"/tasks/{report['id']}", json={"task": "a", "execution_status": "completed"}).status_code == 403


def test_confirm_is_atomic_acknowledged_and_idempotent(org, ai):
    sec, d = _processed_meeting(org, ai)
    mid = d["id"]
    r = sec.post(f"/meetings/{mid}/confirm", json={"draft_revision": d["draft_revision"] - 1})
    assert r.json()["error"]["code"] == "DRAFT_REVISION_MISMATCH"
    assert r.json()["error"]["details"]["current_draft_revision"] == d["draft_revision"]
    r = sec.post(f"/meetings/{mid}/confirm", json={"draft_revision": d["draft_revision"]})
    err = r.json()["error"]
    assert err["code"] == "CONFIRMATION_REQUIRES_ACKNOWLEDGEMENT"
    assert err["details"]["unresolved_speakers"] == ["SPEAKER_02"] and len(err["details"]["tasks_without_assignee"]) == 1
    assert _count("task_assigned") == 0
    # participant cannot confirm
    assert org["api"]["dana"].post(f"/meetings/{mid}/confirm",
                                   json={"draft_revision": d["draft_revision"], "acknowledge_incomplete": True}
                                   ).status_code == 403
    r = sec.post(f"/meetings/{mid}/confirm", json={"draft_revision": d["draft_revision"], "acknowledge_incomplete": True})
    assert r.status_code == 200, r.text
    c = r.json()
    assert c["approval_status"] == "confirmed" and c["protocol_version"] == 1
    assert all(t["status"] == "confirmed" for t in c["tasks"])
    assert _count("task_assigned") == 1  # only the task with an assignee who has an account
    again = sec.post(f"/meetings/{mid}/confirm", json={"draft_revision": c["draft_revision"],
                                                        "acknowledge_incomplete": True})
    assert again.json()["error"]["code"] == "ALREADY_CONFIRMED"
    assert _count("task_assigned") == 1
    # confirmed protocol is frozen
    assert sec.patch(f"/meetings/{mid}/summary", json={"summary": "x"}).json()["error"]["code"] == "MEETING_CONFIRMED"
    assert sec.patch(f"/tasks/{c['tasks'][0]['id']}", json={"task": "x"}).json()["error"]["code"] == "MEETING_CONFIRMED"
    # assignee: sees content and her confirmed task, can complete it; others cannot
    dana, bolat = org["api"]["dana"], org["api"]["bolat"]
    assert dana.get(f"/meetings/{mid}").json()["content_visible"] is True
    mine = dana.get("/tasks?assignee=me").json()
    assert mine["total"] == 1 and mine["items"][0]["task"] == "Подготовить отчёт"
    assert len([n for n in dana.get("/notifications").json()["items"] if n["event_type"] == "task_assigned"]) == 1
    tid = mine["items"][0]["id"]
    assert bolat.patch(f"/tasks/{tid}", json={"execution_status": "completed"}).status_code in (403, 404)
    assert dana.patch(f"/tasks/{tid}", json={"task": "rename"}).status_code == 403
    r = dana.patch(f"/tasks/{tid}", json={"execution_status": "completed"})
    assert r.status_code == 200 and r.json()["execution_status"] == "completed" and r.json()["completed_at"]
    assert dana.get("/tasks?assignee=me&execution_status=completed").json()["total"] == 1
    assert dana.get(f"/tasks?assignee={org['dana']['id']}").status_code == 403


def test_reminders_deduplicated_and_skip_drafts_and_completed(org, ai):
    sec, d = _processed_meeting(org, ai)
    mid = d["id"]
    # deadline 2026-09-25 end of day Almaty = 18:59:59Z
    with db_session() as db:
        early = datetime(2026, 9, 24, 12, 0, tzinfo=timezone.utc)
        assert run_reminders(db, early, 24) == 0  # draft tasks never remind
    r = sec.post(f"/meetings/{mid}/confirm", json={"draft_revision": d["draft_revision"], "acknowledge_incomplete": True})
    assert r.status_code == 200, r.text
    with db_session() as db:
        assert run_reminders(db, datetime(2026, 9, 23, 12, 0, tzinfo=timezone.utc), 24) == 0  # too early
        soon = datetime(2026, 9, 25, 0, 0, tzinfo=timezone.utc)
        created = run_reminders(db, soon, 24)
        assert created == 2  # assignee (dana) + organizer (sec)
        assert run_reminders(db, soon + timedelta(minutes=1), 24) == 0  # repeated tick
        late = datetime(2026, 9, 26, 0, 0, tzinfo=timezone.utc)
        assert run_reminders(db, late, 24) == 2
        assert run_reminders(db, late, 24) == 0
    assert _count("deadline_soon") == 2 and _count("task_overdue") == 2
    dana = org["api"]["dana"]
    tid = dana.get("/tasks?assignee=me").json()["items"][0]["id"]
    assert dana.get("/tasks?assignee=me&overdue=true").json()["total"] in (0, 1)  # depends on real clock
    dana.patch(f"/tasks/{tid}", json={"execution_status": "completed"})
    with db_session() as db:
        assert run_reminders(db, datetime(2026, 10, 5, tzinfo=timezone.utc), 24) == 0


def test_notifications_are_private(org, ai):
    create_meeting(org["api"]["sec"], [org["dana"]["id"]])
    dana, bolat = org["api"]["dana"], org["api"]["bolat"]
    page = dana.get("/notifications").json()
    assert page["unread_count"] == 1
    nid = page["items"][0]["id"]
    assert bolat.patch(f"/notifications/{nid}", json={"read": True}).status_code == 404
    assert dana.patch(f"/notifications/{nid}", json={"read": True}).json()["read_at"] is not None
    assert dana.get("/notifications?unread_only=true").json()["total"] == 0
    dana.patch(f"/notifications/{nid}", json={"read": False})
    assert dana.post("/notifications/read-all").json() == {"updated": 1}


def test_export_uses_saved_edits(org, ai):
    sec, d = _processed_meeting(org, ai)
    mid = d["id"]
    sec.patch(f"/meetings/{mid}/summary", json={"summary": "Қорытынды: есеп дайындалады. Ғ Ү Ұ Қ Ө Һ І"})
    r = sec.get(f"/meetings/{mid}/export?fmt=docx&include_transcript=true")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "_draft" in r.headers["content-disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "ЧЕРНОВИК" in text and "Қорытынды: есеп дайындалады. Ғ Ү Ұ Қ Ө Һ І" in text
    assert "Жақсы, жасаймын." in text  # transcript appendix
    sec.post(f"/meetings/{mid}/confirm", json={"draft_revision": d["draft_revision"] + 1,
                                               "acknowledge_incomplete": True})
    r = org["api"]["dana"].get(f"/meetings/{mid}/export?fmt=docx")
    assert r.status_code == 200 and "_draft" not in r.headers["content-disposition"]
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ЧЕРНОВИК" not in text and "Протокол подтверждён" in text
    cells = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert "Подготовить отчёт" in cells and "Ахметова Дана" in cells


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed locally")
def test_export_pdf(org, ai):
    sec, d = _processed_meeting(org, ai)
    r = sec.get(f"/meetings/{d['id']}/export?fmt=pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_pdf_without_libreoffice_is_explicit_error(org, ai, monkeypatch):
    from app import export as export_mod

    monkeypatch.setattr(export_mod.shutil, "which", lambda _: None)
    sec, d = _processed_meeting(org, ai)
    r = sec.get(f"/meetings/{d['id']}/export?fmt=pdf")
    assert r.status_code == 503 and r.json()["error"]["code"] == "EXPORT_UNAVAILABLE"
